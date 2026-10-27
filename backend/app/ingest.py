"""File ingestion and schema discovery.

Responsibilities:
* validate the uploaded file by extension AND by actual content (magic bytes,
  binary sniffing, decodability),
* read CSV in chunks (no full-file copy just to profile it) and XLSX row by row,
* profile columns (detected type, missingness, samples) and build the 20-row preview,
* suggest a canonical mapping with a confidence and a human-readable reason.

Every value is read as text so nothing is coerced before the user confirms the
mapping; typed parsing happens later in ``cleaning.py``.
"""
from __future__ import annotations

import csv
import math
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterator

import numpy as np
import pandas as pd

from .schemas import CANONICAL_FIELDS

CSV_CHUNK_ROWS = 100_000
PROFILE_SAMPLE_ROWS = 5_000
PREVIEW_ROWS = 20
SUPPORTED_EXTENSIONS = {".csv": "csv", ".xlsx": "xlsx"}


class IngestError(Exception):
    """A user-facing problem with the uploaded file (message is safe to show)."""


@dataclass
class FileFormat:
    kind: str                      # "csv" | "xlsx"
    encoding: str | None = None    # csv only
    delimiter: str | None = None   # csv only
    sheet: str | None = None       # xlsx only
    warnings: list[str] = field(default_factory=list)


# ---------------------------------------------------------------------------
# Format detection
# ---------------------------------------------------------------------------
def detect_format(path: Path, filename: str) -> FileFormat:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise IngestError(f"Unsupported file type '{ext or 'none'}'. Please upload a .csv or .xlsx file.")
    if path.stat().st_size == 0:
        raise IngestError("The file is empty. Please choose a file that contains a header row and data.")

    with open(path, "rb") as fh:
        head = fh.read(65536)

    is_zip = head.startswith(b"PK\x03\x04")
    if SUPPORTED_EXTENSIONS[ext] == "xlsx":
        if not is_zip:
            raise IngestError("This file has an .xlsx extension but is not a valid Excel workbook. "
                              "If it is a CSV, rename it to .csv; otherwise re-save it from Excel as .xlsx.")
        return FileFormat(kind="xlsx")

    # CSV: reject binary content (zip, images, executables...) disguised as .csv
    if is_zip or b"\x00" in head:
        raise IngestError("This file has a .csv extension but contains binary data. "
                          "If it is an Excel workbook, upload it with the .xlsx extension.")
    encoding, warnings = _detect_encoding(path)
    delimiter = _sniff_delimiter(head.decode(encoding, errors="replace"))
    return FileFormat(kind="csv", encoding=encoding, delimiter=delimiter, warnings=warnings)


def _detect_encoding(path: Path) -> tuple[str, list[str]]:
    """UTF-8 is expected. We verify the *whole* file decodes (streamed in 1 MB
    blocks) so a bad byte deep in a large file is caught before import."""
    import codecs
    decoder = codecs.getincrementaldecoder("utf-8-sig")()
    try:
        with open(path, "rb") as fh:
            while block := fh.read(1 << 20):
                decoder.decode(block)
            decoder.decode(b"", final=True)
        return "utf-8-sig", []
    except UnicodeDecodeError:
        return "cp1252", ["The file is not valid UTF-8; it was read as Windows-1252 (Western European). "
                          "Check accented characters in the preview. Re-save as UTF-8 CSV if text looks wrong."]


def _sniff_delimiter(sample: str) -> str:
    lines = sample.splitlines()[:50]
    if not lines:
        return ","
    try:
        dialect = csv.Sniffer().sniff("\n".join(lines), delimiters=",;\t|")
        return dialect.delimiter
    except csv.Error:
        return ","


# ---------------------------------------------------------------------------
# Readers (all values returned as text; empty cells -> NaN)
# ---------------------------------------------------------------------------
def _csv_reader(path: Path, fmt: FileFormat, chunksize: int | None):
    return pd.read_csv(
        path, sep=fmt.delimiter, encoding=fmt.encoding, dtype=str,
        keep_default_na=False, na_values=[""],   # only truly empty cells are null; "NA"/"None" stay text
        skip_blank_lines=False,                  # blank lines are counted, never silently skipped
        chunksize=chunksize, engine="c", low_memory=False,
    )


def _friendly_parser_error(exc: Exception) -> IngestError:
    msg = str(exc)
    m = re.search(r"Expected (\d+) fields in line (\d+), saw (\d+)", msg)
    if m:
        return IngestError(f"The CSV is malformed: line {m.group(2)} has {m.group(3)} fields but the header has "
                           f"{m.group(1)}. This is usually an unquoted comma inside a value. Fix that line "
                           "(or quote the value) and upload again. No rows were imported.")
    if "EOF inside string" in msg or "EOF" in msg:
        return IngestError("The CSV is malformed: a quoted value is never closed. Check for a stray quote "
                           "character and upload again. No rows were imported.")
    return IngestError("The CSV could not be parsed. Check that it is a comma-, semicolon- or tab-separated "
                       "file with one header row. No rows were imported.")


def iter_csv_chunks(path: Path, fmt: FileFormat, chunksize: int = CSV_CHUNK_ROWS) -> Iterator[pd.DataFrame]:
    try:
        reader = _csv_reader(path, fmt, chunksize)
        for chunk in reader:
            yield chunk
    except pd.errors.EmptyDataError:
        raise IngestError("The file has no columns. Please upload a CSV with a header row.")
    except pd.errors.ParserError as exc:
        raise _friendly_parser_error(exc)


def _xlsx_cell_to_text(value) -> str | float:
    if value is None:
        return np.nan
    if isinstance(value, bool):
        return "True" if value else "False"
    if isinstance(value, float):
        if math.isnan(value):
            return np.nan
        return str(int(value)) if value.is_integer() else repr(value)
    if hasattr(value, "isoformat"):
        return value.isoformat()
    text = str(value)
    return text if text.strip() != "" else np.nan


def iter_xlsx_rows(path: Path, fmt: FileFormat) -> tuple[list[str], Iterator[list]]:
    """Open the first non-empty worksheet in read-only (streaming) mode."""
    from openpyxl import load_workbook
    from zipfile import BadZipFile
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except (BadZipFile, KeyError, OSError, ValueError) as exc:
        raise IngestError("The Excel workbook could not be opened. It may be corrupted, password-protected, "
                          "or an older .xls file renamed to .xlsx. Re-save it as .xlsx and try again.") from exc
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        header = None
        for row in rows:
            if row is not None and any(c is not None and str(c).strip() != "" for c in row):
                header = row
                break
        if header is None:
            continue
        fmt.sheet = ws.title
        columns = _dedupe_headers([("" if c is None else str(c).strip()) for c in header])
        width = len(columns)

        def gen():
            for r in rows:
                vals = [_xlsx_cell_to_text(v) for v in (r or ())][:width]
                vals += [np.nan] * (width - len(vals))
                yield vals
            wb.close()
        if len(wb.worksheets) > 1:
            fmt.warnings.append(f"The workbook has {len(wb.worksheets)} sheets; only '{ws.title}' (the first "
                                "sheet with data) was read.")
        return columns, gen()
    wb.close()
    raise IngestError("The workbook contains no data. Please upload a sheet with a header row and data rows.")


def _dedupe_headers(cols: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    out = []
    for i, c in enumerate(cols):
        c = c or f"column_{i + 1}"
        if c in seen:
            seen[c] += 1
            c = f"{c}_{seen[c]}"
        else:
            seen[c] = 0
        out.append(c)
    return out


def read_full(path: Path, fmt: FileFormat) -> pd.DataFrame:
    """Read the whole file as text. CSV is read in chunks and concatenated once."""
    if fmt.kind == "csv":
        chunks = list(iter_csv_chunks(path, fmt))
        df = pd.concat(chunks, ignore_index=True) if chunks else pd.DataFrame()
        df.columns = _dedupe_headers([str(c).strip() for c in df.columns])
        return df
    columns, rows = iter_xlsx_rows(path, fmt)
    return pd.DataFrame(list(rows), columns=columns, dtype=object)


# ---------------------------------------------------------------------------
# Profiling
# ---------------------------------------------------------------------------
BOOL_TOKENS = {"yes", "no", "y", "n", "true", "false", "t", "f", "1", "0"}
RATING_PATTERN = re.compile(r"^\s*-?\d+(\.\d+)?\s*/\s*\d+(\.\d+)?\s*$")


def _to_number(s: pd.Series) -> pd.Series:
    return pd.to_numeric(s.astype(str).str.replace(",", "", regex=False).str.strip(), errors="coerce")


def detect_type(values: pd.Series) -> str:
    """Detected type from a sample of non-null text values."""
    v = values.dropna().astype(str).str.strip()
    v = v[v != ""]
    if v.empty:
        return "empty"
    lower = v.str.lower()
    if lower.isin(BOOL_TOKENS).mean() >= 0.95 and lower.nunique() <= 4 and not lower.isin({"1", "0"}).all():
        return "boolean"
    if v.str.match(RATING_PATTERN).mean() >= 0.5:
        return "rating text (x/y)"
    num = _to_number(v)
    if num.notna().mean() >= 0.95:
        return "integer" if (num.dropna() % 1 == 0).all() else "decimal"
    if num.notna().mean() >= 0.6:
        return "mostly numeric"
    if v.str.contains(r"\d{4}-\d{2}-\d{2}|\d{1,2}/\d{1,2}/\d{2,4}", regex=True).mean() >= 0.9:
        return "date"
    if v.str.contains(",", regex=False).mean() >= 0.3 and v.str.len().median() < 120:
        return "text list"
    return "text"


def profile_file(path: Path, filename: str) -> dict:
    """Stream through the file once: counts, missingness, preview and a sample."""
    fmt = detect_format(path, filename)
    row_count = 0
    null_counts: pd.Series | None = None
    blank_rows = 0
    sample_parts: list[pd.DataFrame] = []
    sample_rows = 0

    if fmt.kind == "csv":
        chunk_iter = iter_csv_chunks(path, fmt)
    else:
        columns, rows = iter_xlsx_rows(path, fmt)

        def xlsx_chunks():
            buf = []
            for r in rows:
                buf.append(r)
                if len(buf) >= CSV_CHUNK_ROWS:
                    yield pd.DataFrame(buf, columns=columns, dtype=object)
                    buf = []
            if buf:
                yield pd.DataFrame(buf, columns=columns, dtype=object)
        chunk_iter = xlsx_chunks()

    columns_out: list[str] | None = None
    for chunk in chunk_iter:
        if columns_out is None:
            columns_out = _dedupe_headers([str(c).strip() for c in chunk.columns])
        chunk.columns = columns_out
        row_count += len(chunk)
        nulls = chunk.isna().sum()
        null_counts = nulls if null_counts is None else null_counts + nulls
        blank_rows += int(chunk.isna().all(axis=1).sum())
        if sample_rows < PROFILE_SAMPLE_ROWS:
            part = chunk.iloc[: PROFILE_SAMPLE_ROWS - sample_rows]
            sample_parts.append(part)
            sample_rows += len(part)

    if columns_out is None:
        # A chunked reader yields nothing for a header-only CSV; distinguish that from "no columns".
        if fmt.kind == "csv":
            try:
                header_only = _csv_reader(path, fmt, None)
                if len(header_only.columns) > 0:
                    raise IngestError("The file has a header row but no data rows. Please upload a file with data.")
            except (pd.errors.EmptyDataError, pd.errors.ParserError):
                pass
        raise IngestError("The file has no readable columns. Please upload a CSV/XLSX with a header row.")
    if row_count == 0 or row_count == blank_rows:
        raise IngestError("The file has a header row but no data rows. Please upload a file with data.")

    sample = pd.concat(sample_parts, ignore_index=True)
    preview = sample.head(PREVIEW_ROWS).astype(object).where(sample.head(PREVIEW_ROWS).notna(), None)
    columns_info = []
    for col in columns_out:
        nc = int(null_counts[col]) if null_counts is not None else 0
        col_sample = sample[col]
        columns_info.append({
            "name": col,
            "detected_type": detect_type(col_sample),
            "null_count": nc,
            "null_pct": round(100 * nc / row_count, 2) if row_count else None,
            "sample_values": [str(x) for x in col_sample.dropna().astype(str).unique()[:5]],
        })

    if len(columns_out) == 1:
        fmt.warnings.append("Only one column was detected. If your file uses a different separator, "
                            "check the preview before importing.")
    mapping = suggest_mapping(columns_out, sample)
    rating_col = next((m["source_column"] for m in mapping if m["suggested_field"] == "rating"), None)
    scale = infer_rating_scale(sample[rating_col]) if rating_col else None
    stated = header_rating_scale(rating_col) if rating_col else None
    if stated:
        scale = {"max": stated, "min": 0.0, "source": f"the header '{rating_col}' states the scale (out of {stated:g})",
                 "confident": True}
    cost_col = next((m["source_column"] for m in mapping if m["suggested_field"] == "cost"), None)
    recognised = {m["suggested_field"] for m in mapping if m["suggested_field"]}
    if len(recognised & RESTAURANT_SIGNAL_FIELDS) < 2:
        fmt.warnings.append("Few restaurant-related columns were recognized. Review the mapping carefully; "
                            "this may not be a restaurant dataset.")

    return {
        "file_type": fmt.kind,
        "encoding": fmt.encoding,
        "delimiter": fmt.delimiter,
        "sheet": fmt.sheet,
        "row_count": row_count,
        "blank_row_count": blank_rows,
        "column_count": len(columns_out),
        "columns": columns_info,
        "preview": preview.to_dict(orient="records"),
        "suggested_mapping": mapping,
        "rating_scale": scale,
        "suggested_price_unit": suggest_price_unit(cost_col) if cost_col else "unknown",
        "warnings": fmt.warnings,
        "profile_sample_rows": len(sample),
    }


# ---------------------------------------------------------------------------
# Mapping suggestions
# ---------------------------------------------------------------------------
RESTAURANT_SIGNAL_FIELDS = {"restaurant_name", "cuisines", "restaurant_type", "rating", "cost",
                            "locality", "city", "review_count", "online_order", "table_booking"}


def normalize_header(name: str) -> str:
    s = re.sub(r"[^0-9a-zA-Z]+", "_", str(name).strip().lower())
    return re.sub(r"_+", "_", s).strip("_")


def _value_check(field_name: str, values: pd.Series) -> tuple[bool, str]:
    """Does the column's content plausibly fit the canonical field?"""
    kind = CANONICAL_FIELDS[field_name]["kind"]
    v = values.dropna().astype(str).str.strip()
    v = v[v != ""]
    if v.empty:
        return True, "column is empty in the sample"
    num = _to_number(v)
    if kind == "rating":
        rating_like = v.str.match(RATING_PATTERN).mean()
        if rating_like >= 0.5:
            return True, "values look like 'x/y' ratings"
        if num.notna().mean() >= 0.5:
            return True, "values are mostly numeric"
        return False, "values are not numeric ratings"
    if kind == "count":
        ok = num.notna().mean() >= 0.8 and (num.dropna() % 1 == 0).mean() >= 0.9
        return ok, "values are whole numbers" if ok else "values are not whole-number counts"
    if kind == "money":
        cleaned = v.str.replace(r"[^\d.\-]", "", regex=True)
        ok = pd.to_numeric(cleaned, errors="coerce").notna().mean() >= 0.7
        return ok, "values are numeric amounts" if ok else "values are not numeric amounts"
    if kind == "bool":
        ok = v.str.lower().isin(BOOL_TOKENS).mean() >= 0.9
        return ok, "values are Yes/No style" if ok else "values are not Yes/No style"
    if kind == "float":
        lim = 90 if field_name == "latitude" else 180
        ok = num.notna().mean() >= 0.9 and (num.dropna().abs() <= lim).all()
        return ok, "values are valid coordinates" if ok else "values are not valid coordinates"
    if kind in {"text", "list"}:
        if num.notna().mean() >= 0.9 and field_name not in {"source_date"}:
            return False, "values are numeric, not text"
        return True, "values are text"
    if kind == "date":
        parsed = pd.to_datetime(v.head(200), errors="coerce", format="mixed")
        ok = parsed.notna().mean() >= 0.8
        return ok, "values parse as dates" if ok else "values do not parse as dates"
    return True, ""


COUNT_TOKENS = {"num", "number", "no", "nb", "count", "counts", "total", "qty", "n"}
ENGAGEMENT_TOKENS = {"rating", "ratings", "review", "reviews", "vote", "votes", "rater", "raters"}


def is_count_header(norm: str) -> bool:
    """'num of ratings', 'number_of_reviews', 'rating count', 'total votes' describe HOW MANY ratings,
    not the rating itself, so they must map to review_count, never to rating."""
    tokens = set(norm.split("_"))
    return bool(tokens & COUNT_TOKENS) and bool(tokens & ENGAGEMENT_TOKENS)


def is_index_column(col: str, values: pd.Series) -> bool:
    """A row index written by pandas/Excel ('Unnamed: 0', 'index', blank header) holding 0..n-1 or 1..n."""
    norm = normalize_header(col)
    if not (norm.startswith("unnamed") or norm in {"index", "column_1", "id_", ""}):
        return False
    num = _to_number(values.dropna()).dropna()
    if num.empty or not (num % 1 == 0).all():
        return False
    diffs = num.diff().dropna()
    return bool(len(diffs) == 0 or (diffs == 1).all())


def suggest_mapping(columns: list[str], sample: pd.DataFrame) -> list[dict]:
    candidates = []  # (score, column, field, reason)
    index_cols = {c for c in columns if is_index_column(c, sample[c])}
    for col in columns:
        if col in index_cols:
            continue
        norm = normalize_header(col)
        tokens = set(norm.split("_"))
        if is_count_header(norm):
            ok, why = _value_check("review_count", sample[col])
            if ok:
                candidates.append((0.93, col, "review_count",
                                   f"header '{col}' describes a count of ratings/reviews; {why}"))
            continue  # never suggest a count column as the rating itself
        for fname, spec in CANONICAL_FIELDS.items():
            best = None
            for idx, syn in enumerate(spec["synonyms"]):
                if norm == syn:
                    best = (0.95 - 0.004 * idx, f"header '{col}' matches '{syn}'")
                    break
                syn_tokens = set(syn.split("_"))
                if len(syn) >= 4 and syn_tokens <= tokens and best is None:
                    best = (0.70 - 0.004 * idx, f"header '{col}' contains '{syn}'")
            if best is None:
                continue
            ok, why = _value_check(fname, sample[col])
            if not ok:
                continue
            candidates.append((best[0], col, fname, f"{best[1]}; {why}"))

    candidates.sort(key=lambda t: (-t[0], columns.index(t[1])))
    assigned_cols: dict[str, tuple] = {}
    assigned_fields: dict[str, str] = {}
    alternatives: dict[str, list[str]] = {}
    for score, col, fname, reason in candidates:
        if col in assigned_cols:
            continue
        if fname in assigned_fields:
            alternatives.setdefault(assigned_fields[fname], []).append(col)
            continue
        assigned_cols[col] = (score, fname, reason)
        assigned_fields[fname] = col

    out = []
    for col in columns:
        if col in assigned_cols:
            score, fname, reason = assigned_cols[col]
            alt = alternatives.get(col)
            if alt:
                reason += f". Other candidate column(s) for this field: {', '.join(alt)} — please confirm."
            out.append({"source_column": col, "suggested_field": fname,
                        "confidence": "high" if score >= 0.9 else "medium",
                        "score": round(score, 3), "reason": reason})
        elif col in index_cols:
            out.append({"source_column": col, "suggested_field": None, "confidence": None, "score": 0,
                        "reason": "Looks like a row index exported by pandas/Excel (0, 1, 2, ...). It is kept as an "
                                  "extra column and can be ignored."})
        else:
            out.append({"source_column": col, "suggested_field": None, "confidence": None, "score": 0,
                        "reason": "No confident match; the column will be preserved as an extra field."})
    return out


def header_rating_scale(col: str) -> float | None:
    """Scale stated in the header itself: 'rate (out of 5)', 'rating/10', 'score_100'."""
    m = re.search(r"(?:out\s*of|/)\s*(\d+(?:\.\d+)?)", str(col), flags=re.I) or \
        re.search(r"(?:^|_)(5|10|100)$", normalize_header(col))
    return float(m.group(1)) if m else None


def infer_rating_scale(values: pd.Series) -> dict:
    """Infer the rating scale from 'x/y' denominators or the observed value range.
    The result is only a suggestion: the user confirms or overrides it at import."""
    v = values.dropna().astype(str).str.strip()
    denoms = v.str.extract(r"/\s*(\d+(?:\.\d+)?)\s*$")[0].dropna()
    if not denoms.empty:
        scale = float(pd.to_numeric(denoms).mode().iloc[0])
        return {"max": scale, "min": 0.0, "source": f"values are written as 'x/{scale:g}'", "confident": True}
    num = _to_number(v).dropna()
    if num.empty:
        return {"max": None, "min": 0.0, "source": "no numeric ratings found", "confident": False}
    p95 = float(num.quantile(0.95))
    for scale in (5.0, 10.0, 100.0):
        if p95 <= scale:
            return {"max": scale, "min": 0.0,
                    "source": f"95% of numeric ratings are <= {scale:g}; please confirm the scale",
                    "confident": False}
    return {"max": None, "min": 0.0, "source": "ratings exceed 100; set the scale manually", "confident": False}


def suggest_price_unit(cost_column: str) -> str:
    norm = normalize_header(cost_column)
    if "two" in norm or "for_2" in norm:
        return "for_two"
    if "person" in norm or "for_one" in norm or "per_head" in norm:
        return "per_person"
    if "range" in norm or "tier" in norm or "level" in norm:
        return "price_tier"
    return "unknown"
